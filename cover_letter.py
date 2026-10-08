from __future__ import annotations

import re
import shutil
import subprocess
from datetime import date
from pathlib import Path

import requests

import config
from matcher import extract_skills, my_projects
from safety import local_ollama_url


def relevant_content(job: dict, profile: dict, resume_info: dict, resume_key: str) -> dict:
    job_text = f"{job.get('title', '')}\n{job.get('required', '')}\n{job.get('preferred', '')}\n{job.get('description', '')}"
    job_skills = extract_skills(job_text)
    mine = set(resume_info.get(resume_key, {}).get("skills", []))
    if not mine:
        mine = set().union(*[set(v.get("skills", [])) for v in resume_info.values()]) if resume_info else set()
    matched = [s for s in sorted(job_skills & mine, key=lambda s: -_importance(s, job))]
    projects = my_projects(profile, resume_info, resume_key)
    words = set(re.findall(r"[a-z]{4,}", job_text.lower()))

    def proj_score(p):
        ps = set(p.get("skills", []))
        pw = set(re.findall(r"[a-z]{4,}", p.get("description", "").lower()))
        return 3 * len(ps & job_skills) + 0.2 * len(pw & words)

    ranked = sorted(projects, key=proj_score, reverse=True)
    return {"matched": matched, "my_skills": mine, "job_skills": job_skills,
            "projects": ([ranked[0]] if ranked else []) + [p for p in ranked[1:2] if proj_score(p) >= 6]}


def _importance(skill: str, job: dict) -> int:
    from matcher import _SKILL_RE
    rx = _SKILL_RE[skill]
    return 3 * len(rx.findall(job.get("required", ""))) + len(rx.findall(job.get("description", "")))


def _join(items: list[str]) -> str:
    items = [i for i in items if i]
    if len(items) <= 1:
        return "".join(items)
    return ", ".join(items[:-1]) + (", and " if len(items) > 2 else " and ") + items[-1]


def _first_sentence(text: str, limit: int = 260) -> str:
    text = re.sub(r"\s+", " ", text or "").strip()
    m = re.match(r"(.+?[.!?])(\s|$)", text)
    s = m.group(1) if m else text
    if len(s) > limit:
        s = s[:limit].rsplit(" ", 1)[0] + "..."
    return s


def _lower_first(s: str) -> str:
    return s[:1].lower() + s[1:] if s and not s[:2].isupper() else s


COMMON_NOUN_SKILLS = {"Machine learning", "Deep learning", "Data structures", "Algorithms", "Statistics",
                      "Data analysis", "Unit testing", "Computer vision", "Prompt engineering", "Vector databases"}
PAST_VERB = re.compile(r"^(?:[A-Z][a-z]+ed|built|wrote|made|led|ran|set|took|grew|drove|designed|developed|"
                       r"created|implemented|trained|deployed|engineered)\b", re.I)


def _nice(skills: list[str]) -> list[str]:
    return [s.lower() if s in COMMON_NOUN_SKILLS else s for s in skills]


def _project_clause(p: dict) -> str:

    first = _first_sentence(p.get("description", ""))
    if PAST_VERB.match(first):
        return f"{p['name']}, where I {_lower_first(first)}"
    return f"{p['name']}. {first}"


def template_letter(job: dict, profile: dict, content: dict, track: str) -> str:
    company, title = job.get("company", "your company"), job.get("title", "internship")
    major, uni, grad = profile.get("major", ""), profile.get("university", ""), profile.get("graduation", "")
    if uni.lower().startswith("university"):
        uni = "the " + uni
    matched = content["matched"]
    projects = content["projects"]
    focus = "building reliable software" if track == "swe" else "building practical machine learning and AI systems"
    paras = [
        f"I am applying for the {title} position at {company}. I am a {major} student at {uni}, "
        f"graduating in {grad}, and I am looking for a summer internship where I can contribute to "
        f"{'production software' if track == 'swe' else 'applied AI/ML work'} while learning from an experienced engineering team."
    ]
    if projects:
        p = projects[0]
        ps = [s for s in p.get("skills", []) if s in content["job_skills"]] or p.get("skills", [])[:3]
        line = f"One project that relates closely to this role is {_project_clause(p)}"
        if ps:
            line += f" It gave me hands-on practice with {_join(_nice(ps[:4]))}, which your posting also asks for." \
                if set(ps) & content["job_skills"] else f" It gave me hands-on practice with {_join(_nice(ps[:4]))}."
        paras.append(line)
    if len(projects) > 1:
        paras.append(f"I also worked on {_project_clause(projects[1])}")
    if matched:
        extra = [s for s in matched if not any(s in pp.get("skills", []) for pp in projects)]
        if extra:
            paras.append(f"Beyond these projects, I have also worked with {_join(_nice(extra[:5]))}, "
                         f"which match tools named in your posting.")
    paras.append(
        f"As an intern on your team, I would bring a habit of {focus}, testing my work, and asking questions early. "
        f"In return, this internship would help me grow in {'software engineering' if track == 'swe' else 'AI/ML engineering'} "
        f"by working on real problems at {company}, which is the direction I want to take after graduation.")
    paras.append(f"Thank you for considering my application. I would welcome the chance to talk about how I can contribute to {company}.")
    return "\n\n".join(paras)


def ollama_available(settings: dict) -> bool:
    if not local_ollama_url(settings.get("ollama_url")):
        return False
    if not settings.get("use_ollama"):
        return False
    try:
        r = requests.get(settings["ollama_url"].rstrip("/") + "/api/tags", timeout=2)
        if r.status_code != 200:
            return False
        names = [m.get("name", "") for m in r.json().get("models", [])]
        want = settings["ollama_model"]
        return any(n == want or n.split(":")[0] == want.split(":")[0] for n in names)
    except requests.RequestException:
        return False


def ollama_letter(job: dict, profile: dict, content: dict, track: str, settings: dict) -> str | None:
    if not local_ollama_url(settings.get("ollama_url")):
        return None
    facts = [f"- Name: {config.full_name(profile)}",
             f"- Student: {profile.get('degree')} in {profile.get('major')}, {profile.get('university')}, graduating {profile.get('graduation')}",
             f"- Skills from resume relevant to this job: {', '.join(content['matched']) or 'none listed'}"]
    for p in content["projects"]:
        facts.append(f"- Project '{p['name']}': {p['description']} (skills: {', '.join(p.get('skills', []))})")
    posting = (job.get("required") or job.get("description") or "")[:2000]
    prompt = (
        "Write the body of a cover letter (4 short paragraphs, 220-330 words, no greeting, no sign-off, no placeholders).\n"
        f"Role: {job.get('title')} at {job.get('company')}.\n"
        "STRICT RULES: Use ONLY the facts below about the candidate. Do not invent skills, tools, numbers, "
        "awards, experience, or facts about the company. Do not use cliches like 'I am thrilled' or 'passionate'. "
        "Plain, natural, confident tone.\n\nCandidate facts:\n" + "\n".join(facts) +
        f"\n\nWhat the posting asks for (for choosing emphasis only):\n{posting}\n")
    try:
        r = requests.post(settings["ollama_url"].rstrip("/") + "/api/generate",
                          json={"model": settings["ollama_model"], "prompt": prompt, "stream": False,
                                "options": {"temperature": 0.4}}, timeout=180)
        r.raise_for_status()
        text = (r.json().get("response") or "").strip()
    except (requests.RequestException, ValueError):
        return None
    return text if letter_is_grounded(text, job, content) else None


def letter_is_grounded(text: str, job: dict, content: dict) -> bool:

    if not text or job.get("company", "").lower() not in text.lower():
        return False
    words = len(text.split())
    if words < 150 or words > 450:
        return False
    if re.search(r"\[(?:your|company|name|insert)[^\]]*\]", text, re.I):
        return False
    claimed = extract_skills(text)
    allowed = content["my_skills"] | {s for p in content["projects"] for s in p.get("skills", [])}
    unsupported = claimed - allowed
    return not unsupported


def safe_name(s: str, n: int = 40) -> str:
    return re.sub(r"[^A-Za-z0-9]+", "_", s or "").strip("_")[:n] or "Unknown"


def write_docx(body: str, job: dict, profile: dict, path: Path) -> None:
    from docx import Document
    from docx.shared import Pt, Inches
    doc = Document()
    for s in doc.sections:
        s.top_margin = s.bottom_margin = Inches(0.9)
        s.left_margin = s.right_margin = Inches(1)
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    style.paragraph_format.space_after = Pt(8)
    name = config.full_name(profile)
    h = doc.add_paragraph()
    r = h.add_run(name)
    r.bold = True
    r.font.size = Pt(14)
    contact = " | ".join(x for x in [profile.get("email"), profile.get("phone"), profile.get("linkedin"),
                                     profile.get("github")] if x)
    if contact:
        doc.add_paragraph(contact)
    doc.add_paragraph(date.today().strftime("%B %d, %Y").replace(" 0", " "))
    doc.add_paragraph(f"Hiring Team\n{job.get('company', '')}")
    doc.add_paragraph("Dear Hiring Team,")
    for para in body.split("\n\n"):
        if para.strip():
            doc.add_paragraph(para.strip())
    doc.add_paragraph(f"Sincerely,\n{name}")
    doc.save(str(path))


def find_soffice() -> str | None:
    for c in ("soffice", "libreoffice"):
        p = shutil.which(c)
        if p:
            return p
    for p in (r"C:\Program Files\LibreOffice\program\soffice.exe",
              r"C:\Program Files (x86)\LibreOffice\program\soffice.exe"):
        if Path(p).exists():
            return p
    return None


def to_pdf(docx_path: Path) -> Path | None:

    soffice = find_soffice()
    if not soffice:
        return None
    try:
        subprocess.run([soffice, "--headless", "--convert-to", "pdf", "--outdir", str(docx_path.parent),
                        str(docx_path)], check=True, capture_output=True, timeout=120)
    except (subprocess.SubprocessError, OSError):
        return None
    pdf = docx_path.with_suffix(".pdf")
    return pdf if pdf.exists() else None


def generate(job: dict, profile: dict, resume_info: dict, resume_key: str, settings: dict,
             want_pdf: bool = True) -> dict:

    config.ensure_dirs()
    content = relevant_content(job, profile, resume_info, resume_key)
    text, method = None, "template"
    if ollama_available(settings):
        print(f"  Writing with local Ollama ({settings['ollama_model']})...")
        text = ollama_letter(job, profile, content, resume_key, settings)
        if text:
            method = "ollama"
        else:
            print("  Ollama's draft didn't pass the accuracy check - using the template instead.")
    if not text:
        text = template_letter(job, profile, content, resume_key)
    stem = f"{safe_name(job.get('company'))}_{safe_name(job.get('title'), 50)}_{date.today():%Y%m%d}_CoverLetter"
    docx_path = config.OUTPUT_DIR / f"{stem}.docx"
    write_docx(text, job, profile, docx_path)
    txt_path = config.OUTPUT_DIR / f"{stem}.txt"
    txt_path.write_text(f"Dear Hiring Team,\n\n{text}\n\nSincerely,\n{config.full_name(profile)}\n", encoding="utf-8")
    pdf_path = to_pdf(docx_path) if want_pdf else None
    return {"docx": str(docx_path), "pdf": str(pdf_path) if pdf_path else "", "txt": str(txt_path),
            "method": method, "text": text}
